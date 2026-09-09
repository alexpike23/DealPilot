#!/usr/bin/env python3
"""Build docs/DealPilot-one-pager.pptx (widescreen single slide)."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x0B, 0x3D, 0x5C)
DARK = RGBColor(0x1A, 0x23, 0x32)
BODY = RGBColor(0x24, 0x30, 0x40)
MUTED = RGBColor(0x5A, 0x65, 0x73)
BG = RGBColor(0xF7, 0xF9, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC5, 0xD0, 0xDB)
OUT = RGBColor(0xD7, 0xDE, 0xE6)

OUT_PATH = "docs/DealPilot-one-pager.pptx"


def _set_run(run, text, size, color, bold=False, name="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name


def _textbox(slide, l, t, w, h, word_wrap=True):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = word_wrap
    tf.auto_size = None
    return tf


def _para(tf, text, size, color, bold=False, space_after=4, first=False, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    run = p.add_run()
    _set_run(run, text, size, color, bold)
    return p


def _runs(tf, parts, size, space_after=4, first=False):
    """parts: list of (text, bold, color)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    for text, bold, color in parts:
        run = p.add_run()
        _set_run(run, text, size, color, bold)
    return p


def _card(slide, l, t, w, h, fill=BG):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = OUT
    sh.line.width = Pt(0.75)
    return sh


def _bar(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def _tight_tf(shape, margin=6):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Pt(margin)
    tf.margin_right = Pt(margin)
    tf.margin_top = Pt(5)
    tf.margin_bottom = Pt(4)
    return tf


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _bar(slide, 0, 0, 13.333, 7.5, WHITE)
    _bar(slide, 0.28, 0.88, 12.77, 0.035, NAVY)

    tf = _textbox(slide, 0.28, 0.16, 9.4, 0.68)
    _para(tf, "DealPilot  —  early-deal copilot for GTM", 22, DARK, True, 2, first=True)
    _para(
        tf,
        "Score the opportunity against how this sales org actually sells, surface leverage for the proposal, then write Salesforce only after Accept.",
        12,
        BODY,
        False,
        0,
    )

    tf = _textbox(slide, 9.7, 0.18, 3.35, 0.62)
    p = _para(tf, "Alex Pike  ·  GTM Business Systems exercise", 11, MUTED, False, 2, first=True, align=PP_ALIGN.RIGHT)
    _para(tf, "Cursor Agent Skill  ·  Salesforce + Apollo", 11, MUTED, False, 0, align=PP_ALIGN.RIGHT)

    # Three columns
    y, h = 1.02, 5.92
    gap = 0.12
    w1, w2, w3 = 3.95, 4.55, 3.95
    x1 = 0.28
    x2 = x1 + w1 + gap
    x3 = x2 + w2 + gap

    _card(slide, x1, y, w1, h)
    _card(slide, x2, y, w2, h)
    _card(slide, x3, y, w3, h)

    # Left
    tf = _textbox(slide, x1 + 0.1, y + 0.08, w1 - 0.2, h - 0.16)
    _para(tf, "THE PROBLEM", 10, NAVY, True, 6, first=True)
    _para(
        tf,
        "Early in a deal, the call, the CRM record, and the pricing playbook rarely get compared. Reps guess tier and discount, copy a competitor’s % onto the wrong SKU, and find out approvals after they have already quoted. Deal Desk then unpicks work that never matched org policy.",
        11,
        BODY,
        False,
        8,
    )
    _runs(
        tf,
        [
            ("What it unlocks: ", True, DARK),
            (
                "a proposal that is legal against internal best practice, with the close levers named (term, packaging, competitive response) and the approval path known before anything is written to Salesforce.",
                False,
                BODY,
            ),
        ],
        11,
        10,
    )
    _para(tf, "ASSUMPTIONS", 10, NAVY, True, 4)
    for item in (
        "qtc_rules.json is org truth (ICP bands, term cap, approval matrix).",
        "ICP revenue comes from Apollo using the Salesforce Account website — never the transcript.",
        "One discovery transcript; seller is the user; they confirm commercial levers.",
        "Dev org has SOAP login, price book Ids, and an Opportunity approval process.",
    ):
        p = _para(tf, "•  " + item, 11, BODY, False, 5)

    # Mid
    tf = _textbox(slide, x2 + 0.1, y + 0.08, w2 - 0.2, 1.35)
    _para(tf, "WHAT I BUILT  &  HOW IT WORKS", 10, NAVY, True, 6, first=True)
    _runs(
        tf,
        [
            ("A Cursor Agent Skill ", True, DARK),
            (
                "(consultative CPQ) plus a Python script that is the only path to Salesforce and Apollo. The model reasons over structured playbook + CRM + transcript; it does not invent revenue or SKUs.",
                False,
                BODY,
            ),
        ],
        11,
        0,
    )

    steps = [
        ("1. Find opp", "Plain-language name search in Salesforce"),
        ("2. Enrich", "Apollo revenue via Account domain"),
        ("3. Score", "Transcript vs playbook; name the levers"),
        ("4. Decide", "Term, then CRM, one at a time"),
        ("5. Accept", "Products, opp fields, approval submit"),
    ]
    sw = 0.78
    sx = x2 + 0.1
    sy = y + 1.52
    sh = 1.18
    gap_s = 0.06
    for i, (title, body) in enumerate(steps):
        box = _card(slide, sx + i * (sw + gap_s), sy, sw, sh, WHITE)
        tf = _tight_tf(box, 4)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        _para(tf, title, 10, NAVY, True, 3, first=True, align=PP_ALIGN.CENTER)
        _para(tf, body, 9, BODY, False, 0, align=PP_ALIGN.CENTER)

    bw = (w2 - 0.28 - 0.08) / 2
    by = sy + sh + 0.12
    bh = h - (by - y) - 0.12
    left_box = _card(slide, x2 + 0.1, by, bw, bh, WHITE)
    tf = _tight_tf(left_box)
    _para(tf, "Systems", 11, NAVY, True, 4, first=True)
    _para(
        tf,
        "Salesforce (read opp/account; write OLIs, Amount, stage, approval comment, submit). Apollo API (X-Api-Key). Local playbook + transcript.",
        11,
        BODY,
        False,
        0,
    )
    right_box = _card(slide, x2 + 0.1 + bw + 0.08, by, bw, bh, WHITE)
    tf = _tight_tf(right_box)
    _para(tf, "Why a skill, not just a script", 11, NAVY, True, 4, first=True)
    _para(
        tf,
        "Script = deterministic I/O. Skill = sequential GTM judgment and an Accept gate so the model cannot commit a quote the seller did not choose.",
        11,
        BODY,
        False,
        0,
    )

    # Right
    tf = _textbox(slide, x3 + 0.1, y + 0.08, w3 - 0.2, h - 0.16)
    _para(tf, "WHERE AI HELPED — AND WHERE I OVERRODE IT", 10, NAVY, True, 6, first=True)
    _para(tf, "•  Cursor sped up Salesforce wiring and the skill scaffold.", 11, BODY, False, 6)
    _runs(
        tf,
        [
            ("•  Did not trust / I changed: ", True, DARK),
            (
                "API key must be X-Api-Key (query-string key 422s). Fail closed if Apollo enrich fails — no invented revenue. No auto multi-year discount. No ZoomInfo 25% on our licenses. No quote + CRM + Accept in one turn.",
                False,
                BODY,
            ),
        ],
        11,
        10,
    )
    _para(tf, "HARDEN FOR PRODUCTION", 10, NAVY, True, 4)
    for item in (
        "OAuth / named credential instead of SOAP password in .env.",
        "Tests on enrich, discount math, and approval routing; logging of every write.",
        "Multi-transcript / multi-opp; human queue proof for the live approval process (submit is in the POC; a real pending approver click is the next proof).",
        "Playbook in Salesforce or a governed store, not only a repo JSON.",
    ):
        _para(tf, "•  " + item, 11, BODY, False, 5)

    _bar(slide, 0.28, 7.05, 12.77, 0.012, LINE)
    tf = _textbox(slide, 0.28, 7.1, 8.2, 0.32)
    _para(
        tf,
        "Left out on purpose: full CPQ UI, e-sign, quoting every SKU. Focused v1: early-deal evaluation + close levers + Salesforce expedite.",
        10,
        MUTED,
        False,
        0,
        first=True,
    )
    tf = _textbox(slide, 8.5, 7.1, 4.55, 0.32)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    _set_run(run, "github.com/alexpike23/DealPilot", 10, NAVY, True)
    run.hyperlink.address = "https://github.com/alexpike23/DealPilot"
    run = p.add_run()
    _set_run(run, "  ·  Demo: “Pike Industries Full Module Deal”", 10, MUTED, False)

    prs.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
